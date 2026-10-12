"""
importador_planilha.py — Importador automático para planilhas do Bolão (.xlsx).

Lê as abas de rodadas (ex: 1ª Rodada, 2ª Rodada...):
  - Cadastra automaticamente os participantes encontrados no cabeçalho
  - Identifica os jogos, times e se é clássico (marcado com estrela ★)
  - Identifica o placar real do jogo (se já preenchido na planilha)
  - Cadastra todos os palpites de cada participante
"""

import re
import unicodedata
import secrets
import openpyxl
from werkzeug.security import generate_password_hash

from database import db
from models import BonusRodada, Jogo, Palpite, Participante, Resultado, Rodada
from services.parser_planilha import parse_placar, extrair_num_rodada, PADRAO_PALPITE, PALAVRAS_IGNORAR


def _normalizar_slug(texto):
    """Gera um slug de email a partir do nome do participante."""
    texto = unicodedata.normalize("NFKD", texto).encode("ascii", "ignore").decode("ascii")
    texto = re.sub(r"[^a-zA-Z0-9]+", ".", texto.strip().lower()).strip(".")
    return texto or "participante"


def importar_planilha_bolao(caminho_ou_arquivo, campeonato_id, *, commit=True, importar_bonus=False):
    """
    Processa arquivo .xlsx e importa rodadas, jogos, participantes e palpites.
    """
    wb = openpyxl.load_workbook(caminho_ou_arquivo, data_only=True)
    try:
        resultado = _importar_workbook(wb, campeonato_id, importar_bonus)
        if commit:
            db.session.commit()
        return resultado
    except Exception:
        db.session.rollback()
        raise
    finally:
        wb.close()


def _importar_workbook(wb, campeonato_id, importar_bonus):

    padrao_palpite = PADRAO_PALPITE
    palavras_ignorar = PALAVRAS_IGNORAR

    total_participantes_novos = 0
    total_rodadas = 0
    total_jogos = 0
    total_palpites = 0

    # Cache de participantes por nome normalizado
    participantes_db = {p.nome.strip().lower(): p for p in Participante.query.all()}

    # Abas que representam rodadas
    abas_rodadas = [s for s in wb.sheetnames if "rodada" in s.lower()]

    for nome_aba in abas_rodadas:
        ws = wb[nome_aba]
        linhas = list(ws.iter_rows(values_only=True))
        if not linhas:
            continue

        nome_rodada_formatado = nome_aba.strip()

        # Localiza ou cria a rodada no campeonato
        rodada = Rodada.query.filter_by(
            campeonato_id=campeonato_id,
            nome=nome_rodada_formatado
        ).first()

        if not rodada:
            rodada = Rodada(nome=nome_rodada_formatado, campeonato_id=campeonato_id, aberta=True)
            db.session.add(rodada)
            db.session.flush()
            total_rodadas += 1

        # Linha 1 tem os participantes a partir da coluna B (índice 1)
        header = linhas[0]
        col_para_participante = {}

        for col_idx in range(1, len(header)):
            val = header[col_idx]
            if val is None:
                continue
            nome_part = str(val).strip()
            if not nome_part or any(p in nome_part.lower() for p in palavras_ignorar):
                continue

            part_key = nome_part.lower()
            if part_key not in participantes_db:
                email_base = _normalizar_slug(nome_part)
                email = f"{email_base}@bolao.local"
                contador = 1
                while Participante.query.filter_by(email=email).first():
                    email = f"{email_base}{contador}@bolao.local"
                    contador += 1

                novo_p = Participante(
                    nome=nome_part,
                    email=email,
                    senha=generate_password_hash(secrets.token_urlsafe(32)),
                    conta_ativa=False,
                    admin=False,
                )
                db.session.add(novo_p)
                db.session.flush()
                participantes_db[part_key] = novo_p
                total_participantes_novos += 1

            col_para_participante[col_idx] = participantes_db[part_key]

        # Linhas de jogos
        for idx in range(1, len(linhas)):
            col0 = linhas[idx][0]
            if not col0:
                continue

            str_col0 = str(col0).strip()
            str_col0_lower = str_col0.lower()

            # Pula linhas de totais, pontuações, cravadas, etc.
            if any(p in str_col0_lower for p in palavras_ignorar):
                continue

            time_casa, time_fora, gols_casa_real, gols_fora_real, especial = parse_placar(str_col0)
            if not time_casa or not time_fora:
                continue
            if gols_casa_real is None and len(linhas[idx]) > 27:
                placar_extra = padrao_palpite.search(str(linhas[idx][27] or ''))
                if placar_extra:
                    gols_casa_real, gols_fora_real = map(int, placar_extra.groups())

            # Localiza ou cria o jogo nessa rodada
            jogo = Jogo.query.filter_by(
                rodada_id=rodada.id,
                time_casa=time_casa,
                time_fora=time_fora
            ).first()

            if not jogo:
                jogo = Jogo(
                    rodada_id=rodada.id,
                    time_casa=time_casa,
                    time_fora=time_fora,
                    especial=especial,
                    encerrado=(gols_casa_real is not None and gols_fora_real is not None),
                )
                db.session.add(jogo)
                db.session.flush()
                total_jogos += 1
            else:
                if especial:
                    jogo.especial = True
                if gols_casa_real is not None and gols_fora_real is not None:
                    jogo.encerrado = True

            # Salva o resultado real se houver
            if gols_casa_real is not None and gols_fora_real is not None:
                resultado = Resultado.query.filter_by(jogo_id=jogo.id).first()
                if not resultado:
                    resultado = Resultado(jogo_id=jogo.id)
                    db.session.add(resultado)
                resultado.gols_casa_real = gols_casa_real
                resultado.gols_fora_real = gols_fora_real

            # Lê os palpites dos participantes nesta linha
            linha_atual = linhas[idx]
            for col_idx, part_obj in col_para_participante.items():
                if col_idx < len(linha_atual):
                    val_palpite = linha_atual[col_idx]
                    if val_palpite is None:
                        continue
                    m_palp = padrao_palpite.search(str(val_palpite))
                    if m_palp:
                        g_casa, g_fora = m_palp.groups()
                        palpite = Palpite.query.filter_by(
                            participante_id=part_obj.id,
                            jogo_id=jogo.id
                        ).first()

                        if not palpite:
                            palpite = Palpite(
                                participante_id=part_obj.id,
                                jogo_id=jogo.id,
                                gols_casa=int(g_casa),
                                gols_fora=int(g_fora),
                            )
                            db.session.add(palpite)
                            total_palpites += 1
                        else:
                            palpite.gols_casa = int(g_casa)
                            palpite.gols_fora = int(g_fora)

    if importar_bonus and 'TOTAL' in wb.sheetnames:
        linhas_total = list(wb['TOTAL'].iter_rows(values_only=True))
        if linhas_total:
            colunas = {i: participantes_db.get(str(nome).strip().lower())
                       for i, nome in enumerate(linhas_total[0]) if i and nome}
            rodadas = {extrair_num_rodada(r.nome): r for r in
                       Rodada.query.filter_by(campeonato_id=campeonato_id).all()}
            iniciou = False
            for linha in linhas_total[1:]:
                titulo = str(linha[0] or '')
                if 'BÔNUS' in titulo.upper() or 'BONUS' in titulo.upper():
                    iniciou = True
                    continue
                if not iniciou:
                    continue
                rodada = rodadas.get(extrair_num_rodada(titulo))
                if rodada is None:
                    continue
                for coluna, pessoa in colunas.items():
                    if pessoa is None or coluna >= len(linha) or linha[coluna] in (None, ''):
                        continue
                    pontos = float(linha[coluna])
                    if not pontos.is_integer() or not 0 <= pontos <= 2147483647:
                        raise ValueError('Bônus inválido na planilha.')
                    bonus = BonusRodada.query.filter_by(participante_id=pessoa.id, rodada_id=rodada.id).first()
                    if bonus is None:
                        bonus = BonusRodada(participante_id=pessoa.id, rodada=rodada, motivo='Bônus histórico da planilha')
                        db.session.add(bonus)
                    bonus.pontos = int(pontos)

    return {
        "participantes_novos": total_participantes_novos,
        "rodadas": total_rodadas,
        "jogos": total_jogos,
        "palpites": total_palpites,
    }
